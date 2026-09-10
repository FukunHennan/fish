from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from pathlib import Path


OUT = Path("/home/chenfukun/fish/docs/机器鱼技术介绍_2026-09-05.pptx")
W, H = 13.333, 7.5

INK = RGBColor(32, 48, 56)
MUTED = RGBColor(101, 119, 123)
TEAL = RGBColor(36, 157, 137)
BLUE = RGBColor(76, 132, 197)
ORANGE = RGBColor(235, 164, 72)
CORAL = RGBColor(210, 104, 113)
MINT = RGBColor(220, 242, 235)
PAPER = RGBColor(250, 252, 250)
LINE = RGBColor(191, 211, 207)
WHITE = RGBColor(255, 255, 255)

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
blank = prs.slide_layouts[6]


def box(slide, x, y, w, h, fill=WHITE, stroke=INK, radius=True, width=1.4):
    s = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = stroke
    s.line.width = Pt(width)
    if radius:
        s.adjustments[0] = 0.13
    return s


def circle(slide, x, y, d, fill=WHITE, stroke=INK, width=1.6):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = stroke
    s.line.width = Pt(width)
    return s


def txt(slide, text, x, y, w, h, size=16, color=INK, bold=False,
        align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    t = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = t.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return t


def arrow(slide, x1, y1, x2, y2, color=TEAL, width=2.0):
    s = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    s.line.color.rgb = color
    s.line.width = Pt(width)
    s.line.end_arrowhead = True
    return s


def title(slide, section, heading, sub=""):
    txt(slide, section.upper(), 0.65, 0.42, 2.8, 0.22, 9, TEAL, True)
    txt(slide, heading, 0.65, 0.78, 11.8, 0.5, 27, INK, True)
    if sub:
        txt(slide, sub, 0.68, 1.38, 11.4, 0.3, 12, MUTED)


def footer(slide, n):
    slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(0.65), Inches(7.08), Inches(12.7), Inches(7.08)
    ).line.color.rgb = LINE
    txt(slide, "FISH CONTROL SYSTEM", 0.65, 7.16, 2.8, 0.18, 8, MUTED, True)
    txt(slide, f"{n:02d}", 12.15, 7.14, 0.55, 0.2, 9, MUTED, True, PP_ALIGN.RIGHT)


def node(slide, x, y, w, h, label, sub, accent=TEAL, fill=WHITE):
    box(slide, x, y, w, h, fill, accent, True, 1.8)
    txt(slide, label, x, y + 0.18, w, 0.28, 16, INK, True, PP_ALIGN.CENTER)
    txt(slide, sub, x + 0.08, y + 0.58, w - 0.16, 0.25, 10, MUTED, False, PP_ALIGN.CENTER)


def note(slide, text, x, y, w, color=TEAL):
    circle(slide, x, y + 0.02, 0.16, color, color, 0.5)
    txt(slide, text, x + 0.28, y, w - 0.28, 0.28, 13, INK)


# 1 - cover
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = PAPER
txt(slide, "FISH CONTROL SYSTEM", 0.8, 0.78, 4, 0.25, 11, TEAL, True)
txt(slide, "机器鱼中央控制系统", 0.8, 1.55, 7.2, 0.72, 34, INK, True)
txt(slide, "设备 + 视觉 + 安全控制", 0.84, 2.52, 5.6, 0.35, 18, MUTED)
txt(slide, "技术介绍 · 2026.09", 0.84, 3.14, 2.1, 0.3, 12, TEAL, True)

# Simple fish icon.
box(slide, 8.65, 1.45, 2.45, 1.25, MINT, TEAL, True, 2.0)
circle(slide, 10.52, 1.82, 0.17, INK, INK, 0.5)
arrow(slide, 8.1, 2.07, 8.65, 2.07, ORANGE, 2.4)
txt(slide, ">", 8.06, 1.88, 0.25, 0.28, 20, ORANGE, True, PP_ALIGN.CENTER)
txt(slide, "鱼", 9.1, 1.77, 0.9, 0.4, 24, TEAL, True, PP_ALIGN.CENTER)
arrow(slide, 11.1, 2.07, 11.75, 2.07, BLUE, 2.4)
txt(slide, "视觉 → 控制 → 运动", 8.0, 3.34, 4.4, 0.3, 13, MUTED, False, PP_ALIGN.CENTER)
footer(slide, 1)

# 2 - architecture
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = PAPER
title(slide, "01 / TOPOLOGY", "系统拓扑：三层，统一入口", "用一张图讲清楚谁负责什么。")
node(slide, 0.95, 2.35, 2.35, 1.25, "浏览器", "React GUI", BLUE, RGBColor(236, 244, 253))
node(slide, 5.05, 2.35, 2.55, 1.25, "Go 控制器", "API / Hub / 权限", TEAL, MINT)
node(slide, 9.6, 1.55, 2.35, 1.25, "Python 视觉", "相机 / YOLO", ORANGE, RGBColor(255, 246, 226))
node(slide, 9.6, 3.55, 2.35, 1.25, "ESP32 机器鱼", "固件 / 舵机 / OTA", CORAL, RGBColor(253, 237, 239))
arrow(slide, 3.3, 2.98, 5.05, 2.98, BLUE)
arrow(slide, 7.6, 2.78, 9.6, 2.18, ORANGE)
arrow(slide, 7.6, 3.18, 9.6, 4.18, CORAL)
txt(slide, "HTTP", 3.7, 2.58, 0.8, 0.22, 10, BLUE, True, PP_ALIGN.CENTER)
txt(slide, "内部接口", 7.85, 2.08, 1.0, 0.22, 10, ORANGE, True, PP_ALIGN.CENTER)
txt(slide, "设备 WebSocket", 7.8, 3.55, 1.3, 0.22, 10, CORAL, True, PP_ALIGN.CENTER)
note(slide, "浏览器不直连 ESP32", 1.08, 5.35, 3.0, CORAL)
note(slide, "Go 负责统一调度", 4.65, 5.35, 3.0, TEAL)
note(slide, "视频和控制分开", 8.2, 5.35, 3.0, BLUE)
footer(slide, 2)

# 3 - control path
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = PAPER
title(slide, "02 / CONTROL", "控制拓扑：一条命令的旅程", "短标签 + 箭头，适合边讲边指。")
items = [
    ("按键", "用户意图", BLUE),
    ("权限", "账号 / clientId", TEAL),
    ("队列", "每鱼独立", ORANGE),
    ("执行", "ESP32 舵机", CORAL),
    ("回报", "ACK / heartbeat", TEAL),
]
xs = [0.75, 3.25, 5.75, 8.25, 10.75]
for i, (h, b, c) in enumerate(items):
    circle(slide, xs[i], 2.55, 1.05, WHITE, c, 2.2)
    txt(slide, h, xs[i], 2.84, 1.05, 0.25, 15, INK, True, PP_ALIGN.CENTER)
    txt(slide, b, xs[i] - 0.25, 3.78, 1.55, 0.25, 11, MUTED, False, PP_ALIGN.CENTER)
    if i < len(items) - 1:
        arrow(slide, xs[i] + 1.08, 3.08, xs[i + 1] - 0.08, 3.08, c, 2.0)
txt(slide, "queued", 2.25, 4.92, 1.2, 0.25, 13, ORANGE, True, PP_ALIGN.CENTER)
txt(slide, "sent", 5.13, 4.92, 1.0, 0.25, 13, BLUE, True, PP_ALIGN.CENTER)
txt(slide, "acknowledged", 7.82, 4.92, 1.7, 0.25, 13, TEAL, True, PP_ALIGN.CENTER)
txt(slide, "控制器收到 ≠ 设备已执行", 4.0, 5.7, 5.3, 0.3, 16, CORAL, True, PP_ALIGN.CENTER)
footer(slide, 3)

# 4 - vision loop
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = PAPER
title(slide, "03 / VISION", "视觉拓扑：识别、计算、回到设备", "视觉身份和物理设备身份必须显式绑定。")
node(slide, 0.85, 2.2, 2.1, 1.2, "摄像头", "最新帧", BLUE, RGBColor(236, 244, 253))
node(slide, 3.45, 2.2, 2.1, 1.2, "YOLO", "检测 / 跟踪", ORANGE, RGBColor(255, 246, 226))
node(slide, 6.05, 2.2, 2.1, 1.2, "路径控制", "坐标 / PID", TEAL, MINT)
node(slide, 8.65, 2.2, 2.1, 1.2, "Go", "会话 / 租约", CORAL, RGBColor(253, 237, 239))
node(slide, 11.05, 2.2, 1.55, 1.2, "鱼", "motion.set", TEAL, WHITE)
for i, c in enumerate([BLUE, ORANGE, TEAL, CORAL]):
    arrow(slide, [2.95, 5.55, 8.15, 10.75][i], 2.8, [3.45, 6.05, 8.65, 11.05][i], 2.8, c)
arrow(slide, 11.8, 3.55, 1.9, 3.55, MUTED, 1.5)
txt(slide, "状态反馈", 5.9, 3.7, 1.2, 0.25, 11, MUTED, True, PP_ALIGN.CENTER)
note(slide, "目标丢失 → 停止", 1.1, 5.0, 2.6, CORAL)
note(slide, "会话过期 → 停止", 4.45, 5.0, 2.6, CORAL)
note(slide, "摄像头失帧 → 停止", 7.8, 5.0, 3.0, CORAL)
footer(slide, 4)

# 5 - safety
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = PAPER
title(slide, "04 / SAFETY", "安全拓扑：四道“停止”防线", "每道防线解决不同类型的失效。")
guards = [
    ("心跳", "设备断线", TEAL),
    ("视觉超时", "Python 卡住", CORAL),
    ("控制权", "窗口接管", BLUE),
    ("参数限幅", "舵机保护", ORANGE),
]
for i, (h, b, c) in enumerate(guards):
    x = 0.92 + i * 3.02
    circle(slide, x + 0.58, 2.18, 1.0, WHITE, c, 2.4)
    txt(slide, "停", x + 0.58, 2.47, 1.0, 0.28, 18, c, True, PP_ALIGN.CENTER)
    txt(slide, h, x, 3.52, 2.16, 0.26, 15, INK, True, PP_ALIGN.CENTER)
    txt(slide, b, x, 3.92, 2.16, 0.24, 11, MUTED, False, PP_ALIGN.CENTER)
    if i < 3:
        arrow(slide, x + 1.62, 2.68, x + 2.74, 2.68, LINE, 1.5)
txt(slide, "任一条件失效，停止命令进入同一设备发送队列", 2.25, 5.28, 8.9, 0.34, 17, CORAL, True, PP_ALIGN.CENTER)
footer(slide, 5)

# 6 - OTA/program
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = PAPER
title(slide, "05 / EXTENSION", "扩展拓扑：OTA 和图形化编程不是一回事", "一个改变能力，一个编排行为，但都经过 Go。")
node(slide, 0.95, 2.22, 2.5, 1.3, "管理员", "上传固件", CORAL, RGBColor(253, 237, 239))
node(slide, 4.08, 2.22, 2.5, 1.3, "学生", "拖拽积木", BLUE, RGBColor(236, 244, 253))
node(slide, 7.25, 2.22, 2.5, 1.3, "Go", "校验 / 调度", TEAL, MINT)
node(slide, 10.35, 2.22, 2.0, 1.3, "ESP32", "能力 / 行为", ORANGE, RGBColor(255, 246, 226))
arrow(slide, 3.45, 2.87, 7.25, 2.87, CORAL)
arrow(slide, 6.58, 2.87, 7.25, 2.87, BLUE)
arrow(slide, 9.75, 2.87, 10.35, 2.87, TEAL)
txt(slide, "OTA：固件分区", 4.35, 3.65, 1.8, 0.25, 12, CORAL, True, PP_ALIGN.CENTER)
txt(slide, "程序：内存运行", 6.55, 3.65, 1.8, 0.25, 12, BLUE, True, PP_ALIGN.CENTER)
note(slide, "升级前先停止", 1.25, 5.15, 2.6, CORAL)
note(slide, "只开放白名单能力", 4.55, 5.15, 3.1, BLUE)
note(slide, "停止 / 急停 / 断线可清除", 8.3, 5.15, 3.6, TEAL)
footer(slide, 6)

# 7 - status
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = PAPER
title(slide, "06 / STATUS", "当前状态：核心链路已成型，实机边界继续验收", "汇报时建议把“已完成”和“待验证”分开说。")
node(slide, 1.0, 2.2, 2.55, 1.32, "Go", "race 测试通过", TEAL, MINT)
node(slide, 4.0, 2.2, 2.55, 1.32, "React", "测试 / 构建通过", BLUE, RGBColor(236, 244, 253))
node(slide, 7.0, 2.2, 2.55, 1.32, "Python", "76 通过 · 1 跳过", ORANGE, RGBColor(255, 246, 226))
node(slide, 10.0, 2.2, 2.25, 1.32, "实机", "停止 ACK 已验证", CORAL, RGBColor(253, 237, 239))
arrow(slide, 3.55, 2.86, 4.0, 2.86, LINE)
arrow(slide, 6.55, 2.86, 7.0, 2.86, LINE)
arrow(slide, 9.55, 2.86, 10.0, 2.86, LINE)
txt(slide, "还要继续测：带载边界、断网停机、双电脑接管、多鱼并发、公网视频", 1.38, 5.22, 10.7, 0.34, 16, CORAL, True, PP_ALIGN.CENTER)
footer(slide, 7)

# 8 - demo flow
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = PAPER
title(slide, "07 / DEMO", "推荐演示顺序：从看见设备，到安全停止", "一条线讲完，观众比较容易跟上。")
demo = [
    ("发现", BLUE),
    ("领取", TEAL),
    ("手动", ORANGE),
    ("视觉", CORAL),
    ("绑定", BLUE),
    ("停止", TEAL),
]
for i, (h, c) in enumerate(demo):
    x = 0.88 + i * 2.03
    circle(slide, x, 2.55, 0.85, WHITE, c, 2.2)
    txt(slide, h, x, 2.82, 0.85, 0.25, 14, INK, True, PP_ALIGN.CENTER)
    if i < len(demo) - 1:
        arrow(slide, x + 0.88, 2.98, x + 1.92, 2.98, c, 2.0)
txt(slide, "讲解主线", 0.95, 4.45, 1.5, 0.25, 15, TEAL, True)
txt(slide, "“Go 是中枢；Python 看环境；ESP32 执行动作；任何失效都回到停止。”", 2.35, 4.42, 9.8, 0.35, 16, INK, True)
txt(slide, "谢谢", 0.8, 5.75, 11.7, 0.55, 28, TEAL, True, PP_ALIGN.CENTER)
footer(slide, 8)

prs.save(OUT)
print(OUT)
