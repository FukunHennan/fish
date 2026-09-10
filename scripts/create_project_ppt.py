from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_AUTO_SIZE
from pathlib import Path


OUT = Path("/home/chenfukun/fish/docs/机器鱼技术介绍_2026-09-05.pptx")

W = 13.333
H = 7.5

NAVY = RGBColor(16, 35, 50)
INK = RGBColor(27, 44, 54)
MUTED = RGBColor(91, 112, 119)
WHITE = RGBColor(250, 252, 250)
PAPER = RGBColor(244, 248, 246)
TEAL = RGBColor(33, 155, 135)
MINT = RGBColor(186, 235, 220)
BLUE = RGBColor(86, 139, 201)
ORANGE = RGBColor(238, 170, 82)
CORAL = RGBColor(218, 112, 120)
LINE = RGBColor(211, 226, 220)
PALE_BLUE = RGBColor(226, 237, 249)
PALE_ORANGE = RGBColor(252, 240, 216)
PALE_CORAL = RGBColor(250, 229, 231)


def rgb(hexstr):
    hexstr = hexstr.lstrip("#")
    return RGBColor(int(hexstr[0:2], 16), int(hexstr[2:4], 16), int(hexstr[4:6], 16))


prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
blank = prs.slide_layouts[6]


def rect(slide, x, y, w, h, fill, line=None, radius=False, transparency=0):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.fill.transparency = transparency
    shape.line.color.rgb = line or fill
    shape.line.width = Pt(0.8)
    if radius:
        shape.adjustments[0] = 0.12
    return shape


def line(slide, x1, y1, x2, y2, color=LINE, width=1.2, dash=None, begin=None, end=None):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    if dash:
        shape.line.dash_style = dash
    if begin:
        shape.line.begin_arrowhead = begin
    if end:
        shape.line.end_arrowhead = end
    return shape


def textbox(slide, text, x, y, w, h, size=16, color=INK, bold=False,
            align=PP_ALIGN.LEFT, font="Microsoft YaHei", valign=MSO_ANCHOR.TOP,
            margin=0.04, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def rich_text(slide, runs, x, y, w, h, size=16, color=INK, align=PP_ALIGN.LEFT,
              margin=0.04, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    for item in runs:
        run = p.add_run()
        run.text = item.get("text", "")
        run.font.name = item.get("font", "Microsoft YaHei")
        run.font.size = Pt(item.get("size", size))
        run.font.bold = item.get("bold", False)
        run.font.color.rgb = item.get("color", color)
    return box


def bullet_list(slide, items, x, y, w, h, size=15, color=INK, gap=0.08, bullet_color=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(gap * 72)
        p.line_spacing = 1.12
        p.alignment = PP_ALIGN.LEFT
        p.text = f"•  {item}"
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(size)
        p.font.color.rgb = color
    return box


def title(slide, kicker, heading, sub=None, dark=False):
    c = WHITE if dark else INK
    mc = MINT if dark else TEAL
    textbox(slide, kicker.upper(), 0.62, 0.42, 3.2, 0.25, 9, mc, True)
    textbox(slide, heading, 0.62, 0.72, 11.7, 0.55, 27, c, True)
    if sub:
        textbox(slide, sub, 0.65, 1.34, 11.5, 0.35, 12, RGBColor(185, 202, 204) if dark else MUTED)


def footer(slide, n, dark=False):
    c = RGBColor(144, 173, 173) if dark else MUTED
    line(slide, 0.62, 7.08, 12.7, 7.08, RGBColor(58, 82, 91) if dark else LINE, 0.8)
    textbox(slide, "FISH CONTROL SYSTEM", 0.62, 7.16, 2.4, 0.18, 8, c, True)
    textbox(slide, f"{n:02d}", 12.22, 7.14, 0.48, 0.2, 9, c, True, PP_ALIGN.RIGHT)


def add_circle(slide, x, y, d, fill, text, text_color=WHITE, size=18):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = fill
    textbox(slide, text, x, y + d * 0.15, d, d * 0.7, size, text_color, True,
            PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    return s


def add_tag(slide, text, x, y, w, fill, color=INK):
    rect(slide, x, y, w, 0.28, fill, fill, True)
    textbox(slide, text, x, y + 0.01, w, 0.2, 9, color, True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def section_card(slide, x, y, w, h, heading, body, accent=TEAL, fill=WHITE, body_size=13):
    rect(slide, x, y, w, h, fill, LINE, True)
    rect(slide, x, y, 0.07, h, accent, accent)
    textbox(slide, heading, x + 0.22, y + 0.2, w - 0.38, 0.3, 15, INK, True)
    textbox(slide, body, x + 0.22, y + 0.62, w - 0.4, h - 0.78, body_size, MUTED)


def flow_box(slide, x, y, w, h, label, sub, fill=WHITE, accent=TEAL, text_color=INK):
    rect(slide, x, y, w, h, fill, accent, True)
    textbox(slide, label, x + 0.08, y + 0.17, w - 0.16, 0.28, 14, text_color, True, PP_ALIGN.CENTER)
    textbox(slide, sub, x + 0.08, y + 0.54, w - 0.16, 0.26, 9.5, MUTED if fill != NAVY else MINT, False, PP_ALIGN.CENTER)


# 1
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = NAVY
rect(slide, 0, 0, W, H, NAVY, NAVY)
rect(slide, 0, 0, 0.22, H, TEAL, TEAL)
textbox(slide, "FISH CONTROL SYSTEM", 0.78, 0.68, 4.0, 0.3, 11, MINT, True)
textbox(slide, "机器鱼中央控制系统", 0.78, 1.46, 7.5, 0.82, 34, WHITE, True)
textbox(slide, "从设备接入、视觉识别到安全运动控制的一体化技术方案", 0.82, 2.45, 7.6, 0.42, 17, RGBColor(202, 220, 219))
add_tag(slide, "技术介绍 · 2026.09", 0.82, 3.18, 1.6, MINT, NAVY)

# stylized fish and signal lines
add_circle(slide, 9.92, 1.18, 1.78, TEAL, "鱼", WHITE, 30)
rect(slide, 9.12, 2.78, 3.38, 0.06, MINT, MINT, True)
for i, (yy, col, txt) in enumerate([(3.18, BLUE, "视觉"), (3.73, ORANGE, "控制"), (4.28, CORAL, "设备")]):
    line(slide, 8.35, yy, 11.85, yy, col, 2.3)
    add_circle(slide, 8.15, yy - 0.11, 0.22, col, "", WHITE, 1)
    textbox(slide, txt, 11.95, yy - 0.12, 0.65, 0.25, 11, col, True)
textbox(slide, "浏览器  →  Go  →  Python / ESP32", 7.62, 5.36, 4.55, 0.32, 12, RGBColor(202, 220, 219), False, PP_ALIGN.CENTER)
footer(slide, 1, True)

# 2
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = PAPER
title(slide, "01 / WHY", "先把问题说清楚：控制一条鱼，实际是管理一套系统",
      "项目的核心不是“让舵机动起来”，而是把设备、视觉、权限、网络和恢复机制放到同一条可验证链路里。")
section_card(slide, 0.72, 2.08, 3.82, 1.62, "设备层", "ESP32 负责运动执行、心跳、设备身份、联网、状态灯和 OTA。", TEAL, WHITE)
section_card(slide, 4.76, 2.08, 3.82, 1.62, "控制层", "Go 是唯一控制入口，统一接入设备、视觉、用户会话和命令队列。", BLUE, WHITE)
section_card(slide, 8.80, 2.08, 3.82, 1.62, "感知层", "Python 负责摄像头、YOLO、跟踪、坐标映射、路径和视觉控制计算。", ORANGE, WHITE)
textbox(slide, "项目要解决的四个技术问题", 0.75, 4.22, 4.4, 0.3, 17, INK, True)
items = [
    ("统一入口", "浏览器不直连 ESP32，所有命令经过 Go。"),
    ("闭环控制", "视频、识别、控制量和设备回报形成闭环。"),
    ("安全边界", "控制权、超时、急停、断线保护分别成立。"),
    ("可演进", "OTA 与 Mixly 风格编程复用现有能力和安全链路。"),
]
for i, (h, b) in enumerate(items):
    x = 0.78 + (i % 2) * 6.08
    y = 4.72 + (i // 2) * 0.78
    add_circle(slide, x, y + 0.02, 0.3, [TEAL, BLUE, CORAL, ORANGE][i], str(i + 1), WHITE, 10)
    textbox(slide, h, x + 0.43, y, 1.3, 0.25, 13, INK, True)
    textbox(slide, b, x + 1.7, y, 4.1, 0.35, 11, MUTED)
footer(slide, 2)

# 3
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = PAPER
title(slide, "02 / ARCHITECTURE", "总体架构：浏览器只面对 Go，Go 负责秩序",
      "三进程、一个设备协议、两条媒体边界：控制面和视频面分开设计。")
flow_box(slide, 0.82, 2.1, 2.25, 1.05, "React GUI", "手动 / 视觉 / 设置", PALE_BLUE, BLUE)
flow_box(slide, 4.08, 2.1, 2.55, 1.05, "Go Controller", "API / Hub / Auth / Queue", NAVY, NAVY, WHITE)
flow_box(slide, 8.0, 1.55, 2.2, 1.05, "Python Vision", "Camera / YOLO / WebRTC", PALE_ORANGE, ORANGE)
flow_box(slide, 8.0, 3.0, 2.2, 1.05, "ESP32 Fish", "Firmware / Motion / OTA", PALE_CORAL, CORAL)
line(slide, 3.08, 2.62, 4.08, 2.62, TEAL, 2.3)
line(slide, 6.63, 2.62, 8.0, 2.08, BLUE, 2.3)
line(slide, 6.63, 2.62, 8.0, 3.52, CORAL, 2.3)
textbox(slide, "HTTP / WebSocket", 3.05, 2.24, 1.15, 0.24, 9, TEAL, True, PP_ALIGN.CENTER)
textbox(slide, "内部 HTTP", 6.78, 1.77, 0.9, 0.24, 9, BLUE, True, PP_ALIGN.CENTER)
textbox(slide, "设备 WebSocket", 6.77, 3.5, 1.12, 0.24, 9, CORAL, True, PP_ALIGN.CENTER)
rect(slide, 0.82, 4.72, 11.7, 1.12, WHITE, LINE, True)
textbox(slide, "关键边界", 1.05, 4.96, 1.2, 0.24, 13, INK, True)
bullet_list(slide, [
    "Python 不直接访问鱼的 IP；视觉运动意图必须经过 Go。",
    "WebRTC 传视频，Go 代理信令与 API，不把视频字节塞进设备 WebSocket。",
    "ESP32 只执行最终运动参数，不接收视频、PID 或网页身份。",
], 2.28, 4.91, 9.8, 0.72, 12, MUTED)
footer(slide, 3)

# 4
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = NAVY
title(slide, "03 / CONTROL PATH", "一次“前进”命令，如何走完端到端链路", "每台设备拥有独立发送队列；状态以设备回报为准。", True)
steps = [
    ("1", "浏览器", "按键 / 视觉意图"),
    ("2", "Go Web", "认证 + 租约 + 参数"),
    ("3", "Hub", "单设备命令队列"),
    ("4", "ESP32", "执行 motion.set"),
    ("5", "反馈", "ACK + heartbeat"),
]
xs = [0.78, 3.18, 5.58, 7.98, 10.38]
for i, (num, h, b) in enumerate(steps):
    add_circle(slide, xs[i], 2.33, 0.55, [BLUE, TEAL, ORANGE, CORAL, MINT][i], num, NAVY if i == 4 else WHITE, 15)
    textbox(slide, h, xs[i] - 0.18, 3.06, 1.6, 0.28, 14, WHITE, True, PP_ALIGN.CENTER)
    textbox(slide, b, xs[i] - 0.3, 3.42, 1.85, 0.44, 11, RGBColor(188, 207, 208), False, PP_ALIGN.CENTER)
    if i < len(steps) - 1:
        line(slide, xs[i] + 0.58, 2.61, xs[i + 1] - 0.08, 2.61, MINT, 1.7)
textbox(slide, "协议层的三个语义", 0.82, 4.62, 2.9, 0.3, 16, MINT, True)
proto = [
    ("queued", "已进入控制器队列，不代表设备已执行"),
    ("sent", "已写入设备 WebSocket"),
    ("acknowledged", "设备返回同 requestId 的结果"),
]
for i, (h, b) in enumerate(proto):
    y = 5.04 + i * 0.47
    textbox(slide, h, 0.95, y, 1.7, 0.25, 12, [ORANGE, BLUE, MINT][i], True)
    textbox(slide, b, 2.75, y, 7.5, 0.25, 11, RGBColor(205, 218, 218))
footer(slide, 4, True)

# 5
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = PAPER
title(slide, "04 / VISION LOOP", "视觉不是“识别完就发命令”，而是一条有身份绑定的闭环",
      "YOLO 负责低频身份确认，标记 / 跟踪负责更高频观测，最终控制量仍回到 Go 的设备队列。")
nodes = [
    (0.72, "Camera", "1920×1080\n目标 30 FPS", PALE_BLUE, BLUE),
    (3.02, "Perception", "YOLO + marker\ntracking", PALE_ORANGE, ORANGE),
    (5.32, "Mapping", "坐标 / 标定\n路径 / 航向", MINT, TEAL),
    (7.62, "Controller", "PID / motion\nfrequency + bias", PALE_CORAL, CORAL),
    (9.92, "Go → Fish", "session + lease\nmotion.set", NAVY, NAVY),
]
for i, (x, h, b, fill, accent) in enumerate(nodes):
    flow_box(slide, x, 2.34, 1.72, 1.16, h, b, fill, accent, WHITE if fill == NAVY else INK)
    if i < len(nodes) - 1:
        line(slide, x + 1.72, 2.92, nodes[i + 1][0], 2.92, accent, 1.8)
textbox(slide, "视觉身份 ≠ 物理设备身份", 0.82, 4.25, 3.6, 0.3, 16, INK, True)
section_card(slide, 0.82, 4.76, 3.72, 1.18, "必须显式绑定", "识别目标 ID 与 ESP32 deviceId 不能靠“看起来像”自动对应。", BLUE, WHITE, 12)
section_card(slide, 4.8, 4.76, 3.72, 1.18, "工作区隔离", "每个账号 + 浏览器 + 鱼拥有独立视觉控制工作区。", TEAL, WHITE, 12)
section_card(slide, 8.78, 4.76, 3.72, 1.18, "失效即停止", "目标丢失、会话过期、摄像头失帧或控制权变化都触发停止。", CORAL, WHITE, 12)
footer(slide, 5)

# 6
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = PAPER
title(slide, "05 / SAFETY", "安全设计：把“停止”做成多层防线",
      "不同超时解决不同问题，不能用一个 60 秒租约替代设备心跳或视觉运动保护。")
rows = [
    ("设备心跳", "约 3 秒", "设备 / Go 断连后停止", TEAL),
    ("视觉运动有效期", "约 3 秒", "Python 卡住或视觉失效时停止", CORAL),
    ("手动控制租约", "默认 60 秒", "控制权归属与接管，不等于持续运动", BLUE),
    ("舵机参数限幅", "中位 ± 振幅", "Go 联合校验，避免超出标定范围", ORANGE),
]
for i, (h, t, b, accent) in enumerate(rows):
    y = 2.03 + i * 0.9
    rect(slide, 0.78, y, 11.72, 0.67, WHITE, LINE, True)
    rect(slide, 0.78, y, 0.09, 0.67, accent, accent)
    textbox(slide, h, 1.05, y + 0.17, 2.0, 0.25, 13, INK, True)
    add_tag(slide, t, 3.12, y + 0.18, 1.12, [MINT, PALE_CORAL, PALE_BLUE, PALE_ORANGE][i], accent)
    textbox(slide, b, 4.62, y + 0.17, 6.9, 0.26, 12, MUTED)
textbox(slide, "认证链路", 0.82, 5.82, 1.3, 0.25, 15, INK, True)
textbox(slide, "网页账号  →  浏览器 clientId  →  设备租约  →  HMAC challenge / response  →  WebSocket 注册", 2.04, 5.82, 10.2, 0.25, 12, MUTED)
footer(slide, 6)

# 7
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = NAVY
title(slide, "06 / MULTI-USER", "多用户控制：账号拥有设备，不代表当前窗口可以发运动",
      "控制归属和控制会话拆开，避免刷新、换窗口或接管后旧命令继续生效。", True)
rect(slide, 0.82, 2.1, 4.1, 3.76, RGBColor(24, 52, 63), RGBColor(54, 86, 92), True)
textbox(slide, "状态转换", 1.08, 2.42, 1.6, 0.28, 16, MINT, True)
states = [("空闲", TEAL), ("领取", BLUE), ("当前窗口控制", ORANGE), ("释放 / 接管", CORAL)]
for i, (s, c) in enumerate(states):
    y = 3.0 + i * 0.62
    add_circle(slide, 1.12, y, 0.33, c, str(i + 1), NAVY if c == ORANGE else WHITE, 10)
    textbox(slide, s, 1.62, y + 0.03, 2.4, 0.24, 13, WHITE, i == 2)
    if i < len(states) - 1:
        line(slide, 1.28, y + 0.34, 1.28, y + 0.62, RGBColor(121, 168, 162), 1.2)
textbox(slide, "旧窗口命令被拒绝\n停止命令仍优先入队", 1.12, 5.22, 3.1, 0.45, 12, RGBColor(202, 220, 219))
textbox(slide, "控制权判断的三个条件", 5.58, 2.22, 3.7, 0.28, 16, MINT, True)
checks = [
    "userId 匹配当前登录账号",
    "clientId 匹配当前浏览器窗口",
    "deviceId 对应当前设备租约",
]
for i, text in enumerate(checks):
    y = 2.92 + i * 0.68
    add_circle(slide, 5.62, y, 0.33, MINT, "✓", NAVY, 13)
    textbox(slide, text, 6.15, y + 0.02, 4.7, 0.25, 13, WHITE, True)
textbox(slide, "工程意义", 5.58, 5.25, 1.6, 0.25, 15, MINT, True)
textbox(slide, "刷新页面可以恢复登录，但不能自动恢复运动会话；恢复控制必须显式发生。", 7.03, 5.23, 4.5, 0.5, 12, RGBColor(202, 220, 219))
footer(slide, 7, True)

# 8
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = PAPER
title(slide, "07 / OTA + PROGRAMMING", "OTA 与图形化编程：两类任务，共用一条受控执行链",
      "固件升级改变设备能力；临时程序编排设备行为。二者都不能绕过 Go 的安全闸门。")
section_card(slide, 0.82, 2.05, 5.55, 3.6, "OTA：改变固件能力", "管理员上传 firmware.bin + manifest\n\n预检查：硬件 / 能力 / 版本 / SHA-256\n\n升级前停止运动 → 写入备用分区 → 重启 → 重新注册确认\n\n失败时保留身份配置，并进入回滚或待核查流程。", CORAL, WHITE, 14)
section_card(slide, 6.95, 2.05, 5.55, 3.6, "Mixly 风格：编排临时行为", "学生程序只调用底层固件的白名单能力\n\n运动、灯光、传感器、条件、循环和计时\n\n程序运行在控制器内存，不写入固件分区\n\n停止、急停、断线或重启后，运行实例自动清除。", TEAL, WHITE, 14)
line(slide, 6.58, 2.3, 6.58, 5.38, LINE, 1.4)
add_tag(slide, "同一控制器", 5.58, 5.92, 2.0, MINT, NAVY)
textbox(slide, "共享：能力描述、设备队列、控制权、运动限幅、超时保护、急停", 3.2, 6.33, 6.95, 0.25, 12, MUTED, False, PP_ALIGN.CENTER)
footer(slide, 8)

# 9
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = PAPER
title(slide, "08 / ENGINEERING", "工程验证：自动化已经覆盖核心逻辑，实机边界仍需诚实标注",
      "截至 2026 年 9 月 4 日维护记录，项目已形成 Go、前端、Python、固件测试的组合验证。")
metrics = [
    ("Go", "go test -race ./...", "通过", TEAL),
    ("React", "npm test / build", "通过", BLUE),
    ("Python", "76 通过 · 1 跳过", "通过", ORANGE),
    ("实机", "停止 ACK / 重连", "部分验证", CORAL),
]
for i, (h, b, s, c) in enumerate(metrics):
    x = 0.82 + i * 3.04
    rect(slide, x, 2.08, 2.62, 1.42, WHITE, LINE, True)
    add_circle(slide, x + 0.22, 2.33, 0.35, c, "", WHITE, 1)
    textbox(slide, h, x + 0.72, 2.31, 1.4, 0.24, 15, INK, True)
    textbox(slide, b, x + 0.22, 2.84, 2.2, 0.24, 10.5, MUTED)
    add_tag(slide, s, x + 0.22, 3.14, 0.88 if s == "通过" else 1.22, MINT if s == "通过" else PALE_ORANGE, c)
textbox(slide, "已验证的能力", 0.82, 4.2, 2.1, 0.28, 16, INK, True)
bullet_list(slide, [
    "视觉超时停止、旧会话拒绝、租约接管和发送队列生命周期。",
    "相机裁剪、WebRTC/API、设备发现、认证、运动限幅和状态恢复。",
    "本机部署健康检查：Go 8081、Python 8091、前端 8098。",
], 0.92, 4.66, 6.0, 1.35, 12, MUTED)
textbox(slide, "尚待实机验收", 7.2, 4.2, 2.1, 0.28, 16, INK, True)
bullet_list(slide, [
    "带载舵机边界与固件中位和 Go 标定的一致性。",
    "运行中断网 / 停 Python 的物理停机时延。",
    "两台真实电脑接管、多鱼并行控制和公网媒体长时间稳定性。",
], 7.3, 4.66, 5.35, 1.35, 12, MUTED)
footer(slide, 9)

# 10
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = NAVY
title(slide, "09 / NEXT", "下一步：从“能控制”走向“可规模化运行”", "建议按依赖顺序推进，先收敛身份和任务模型，再扩展并发与生态。", True)
roadmap = [
    ("现在", "稳定控制基础", "单设备队列、租约、视觉超时、OTA 基础链路", TEAL),
    ("下一阶段", "多鱼视觉隔离", "每鱼目标绑定、独立工作区、稳定目标身份", BLUE),
    ("再下一步", "任务与生态", "OTA 调度、能力目录、Mixly 程序解释器、审计", ORANGE),
    ("验收重点", "真实运行指标", "物理停机、并发、断网、公网媒体、长时间运行", CORAL),
]
for i, (when, h, b, c) in enumerate(roadmap):
    x = 0.82 + i * 3.05
    rect(slide, x, 2.14, 2.64, 2.68, RGBColor(24, 52, 63), RGBColor(54, 86, 92), True)
    add_tag(slide, when, x + 0.22, 2.4, 0.82, c, NAVY if c in [MINT, ORANGE] else WHITE)
    textbox(slide, h, x + 0.22, 2.98, 2.15, 0.34, 15, WHITE, True)
    textbox(slide, b, x + 0.22, 3.58, 2.15, 0.7, 11, RGBColor(202, 220, 219))
    if i < len(roadmap) - 1:
        line(slide, x + 2.64, 3.48, x + 3.05, 3.48, MINT, 1.6)
textbox(slide, "演示建议", 0.82, 5.42, 1.45, 0.28, 16, MINT, True)
textbox(slide, "设备发现 → 领取控制权 → 手动动作 → 启动视觉 → 目标绑定 → 视觉停止保护 → OTA 预检查", 2.22, 5.42, 9.95, 0.3, 13, WHITE, True)
textbox(slide, "一句话总结：Go 把设备、视觉和用户控制收敛为一条可验证、可恢复、可扩展的控制链路。", 0.82, 6.18, 11.7, 0.32, 15, MINT, True, PP_ALIGN.CENTER)
footer(slide, 10, True)

prs.save(OUT)
print(OUT)
