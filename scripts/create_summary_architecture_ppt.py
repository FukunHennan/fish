from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path("/home/chenfukun/fish/docs/机器鱼技术介绍_2026-09-05.pptx")
W, H = 13.333, 7.5

BLACK = RGBColor(30, 37, 40)
GRAY = RGBColor(232, 232, 232)
GRAY2 = RGBColor(246, 246, 246)
GRAY3 = RGBColor(170, 180, 181)
WHITE = RGBColor(255, 255, 255)
TEAL = RGBColor(30, 137, 119)
BLUE = RGBColor(61, 106, 158)
ORANGE = RGBColor(179, 113, 31)
RED = RGBColor(180, 72, 82)
MINT = RGBColor(225, 243, 237)
PALE_BLUE = RGBColor(238, 245, 253)
PALE_ORANGE = RGBColor(255, 247, 230)
PALE_RED = RGBColor(253, 239, 241)

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
blank = prs.slide_layouts[6]


def text(slide, value, x, y, w, h, size=14, color=BLACK, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = s.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.035)
    tf.margin_right = Inches(0.035)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = value
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return s


def rounded(slide, x, y, w, h, fill=WHITE, stroke=BLACK, lw=1.1):
    s = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = stroke
    s.line.width = Pt(lw)
    s.adjustments[0] = 0.08
    return s


def line(slide, x1, y1, x2, y2, color=BLACK, width=1.2, arrow=False):
    s = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    s.line.color.rgb = color
    s.line.width = Pt(width)
    if arrow:
        s.line.end_arrowhead = True
    return s


def module(slide, x, y, w, h, label, fill=GRAY, size=11, color=BLACK):
    rounded(slide, x, y, w, h, fill, BLACK, 0.8)
    text(slide, label, x + 0.03, y + 0.02, w - 0.06, h - 0.04, size, color, False, PP_ALIGN.CENTER)


def group(slide, x, y, w, h, label, accent):
    rounded(slide, x, y, w, h, WHITE, BLACK, 1.2)
    text(slide, label, x + 0.1, y + 0.05, w - 0.2, 0.25, 13, accent, True, PP_ALIGN.CENTER)


def heading(slide, kicker, title, sub):
    text(slide, kicker, 0.62, 0.3, 3.5, 0.2, 9, TEAL, True)
    text(slide, title, 0.62, 0.59, 12, 0.45, 24, BLACK, True)
    text(slide, sub, 0.64, 1.08, 11.8, 0.25, 11, GRAY3)


def footer(slide, number):
    line(slide, 0.56, 7.08, 12.78, 7.08, GRAY3, 0.7)
    text(slide, "FISH CONTROL SYSTEM", 0.58, 7.14, 2.5, 0.18, 8, GRAY3, True)
    text(slide, f"{number:02d}", 12.15, 7.13, 0.55, 0.2, 9, GRAY3, True, PP_ALIGN.RIGHT)


# 1. What is the system?
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = WHITE
heading(slide, "01 / WHAT", "项目是什么：一套机器鱼控制中枢", "核心价值：把设备、视觉和用户操作统一到一条链路。")

rounded(slide, 0.55, 1.52, 12.22, 5.25, WHITE, BLACK, 1.3)
text(slide, "机器鱼中央控制系统", 0.78, 1.64, 3.2, 0.25, 13, BLACK, True)

# Read top-to-bottom: user intent -> central orchestration -> vision/device execution.
group(slide, 3.82, 1.98, 5.68, 0.95, "用户端", BLUE)
module(slide, 4.25, 2.3, 1.45, 0.4, "React GUI", PALE_BLUE, 10)
module(slide, 5.95, 2.3, 1.45, 0.4, "登录 / 选择", PALE_BLUE, 10)
module(slide, 7.65, 2.3, 1.45, 0.4, "手动 / 视觉", PALE_BLUE, 10)

group(slide, 3.82, 3.15, 5.68, 0.95, "Go 控制中枢", TEAL)
module(slide, 4.25, 3.47, 1.45, 0.4, "API", MINT, 10)
module(slide, 5.95, 3.47, 1.45, 0.4, "权限 / 租约", MINT, 10)
module(slide, 7.65, 3.47, 1.45, 0.4, "设备队列", MINT, 10)

group(slide, 1.12, 4.35, 5.68, 1.72, "Python 视觉", ORANGE)
module(slide, 1.52, 4.8, 1.45, 0.48, "相机", PALE_ORANGE, 11)
module(slide, 3.25, 4.8, 1.45, 0.48, "YOLO", PALE_ORANGE, 11)
module(slide, 4.98, 4.8, 1.45, 0.48, "路径 / PID", PALE_ORANGE, 10)
module(slide, 2.38, 5.45, 1.45, 0.38, "跟踪 / 标定", GRAY2, 9)
module(slide, 4.1, 5.45, 1.45, 0.38, "WebRTC", GRAY2, 9)

group(slide, 7.25, 4.35, 5.35, 1.72, "ESP32 机器鱼", RED)
module(slide, 7.62, 4.8, 1.35, 0.48, "认证", PALE_RED, 11)
module(slide, 9.22, 4.8, 1.35, 0.48, "运动", PALE_RED, 11)
module(slide, 10.82, 4.8, 1.35, 0.48, "心跳", PALE_RED, 11)
module(slide, 8.42, 5.45, 1.35, 0.38, "WebSocket", GRAY2, 9)
module(slide, 10.02, 5.45, 1.35, 0.38, "OTA", GRAY2, 9)

line(slide, 6.66, 2.93, 6.66, 3.15, BLUE, 1.8, True)
line(slide, 6.66, 4.1, 6.66, 4.35, TEAL, 1.8, True)
line(slide, 5.95, 4.1, 3.96, 4.35, ORANGE, 1.5, True)
line(slide, 7.38, 4.1, 9.9, 4.35, RED, 1.5, True)
text(slide, "意图", 6.82, 2.98, 0.5, 0.18, 9, BLUE, True)
text(slide, "调度", 6.82, 4.14, 0.5, 0.18, 9, TEAL, True)

footer(slide, 1)


# 2. How does it work?
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = WHITE
heading(slide, "02 / HOW", "系统怎么工作：感知和控制形成闭环", "关键不是单向发命令，而是“识别 → 控制 → 执行 → 反馈”。")

rounded(slide, 0.62, 1.55, 12.1, 4.95, WHITE, BLACK, 1.3)
text(slide, "视觉控制链路", 0.9, 1.75, 1.8, 0.25, 13, ORANGE, True)

items = [
    ("摄像头", "采集画面", BLUE, PALE_BLUE),
    ("YOLO", "识别 / 跟踪", ORANGE, PALE_ORANGE),
    ("路径控制", "坐标 / PID", TEAL, MINT),
    ("Go 控制器", "权限 / 队列", RED, PALE_RED),
    ("ESP32", "运动执行", TEAL, WHITE),
]
xs = [0.92, 3.25, 5.58, 7.91, 10.24]
for i, (h, b, c, fill) in enumerate(items):
    module(slide, xs[i], 2.45, 1.85, 0.92, h, fill, 14)
    text(slide, b, xs[i], 3.58, 1.85, 0.22, 10, GRAY3, False, PP_ALIGN.CENTER)
    if i < len(items) - 1:
        line(slide, xs[i] + 1.87, 2.91, xs[i + 1], 2.91, c, 1.8, True)

line(slide, 11.15, 3.95, 1.85, 3.95, GRAY3, 1.2, True)
text(slide, "状态反馈：heartbeat / state / command.result", 4.15, 4.03, 5.0, 0.25, 11, GRAY3, True, PP_ALIGN.CENTER)

rounded(slide, 1.0, 4.78, 5.25, 0.95, PALE_RED, RED, 0.8)
text(slide, "身份绑定", 1.25, 4.97, 1.0, 0.23, 13, RED, True)
text(slide, "视觉目标 ID ≠ 物理 deviceId，必须显式绑定。", 2.38, 4.93, 3.5, 0.3, 12, BLACK)

rounded(slide, 7.02, 4.78, 5.25, 0.95, MINT, TEAL, 0.8)
text(slide, "停止条件", 7.28, 4.97, 1.0, 0.23, 13, TEAL, True)
text(slide, "目标丢失、会话过期、失帧或断线，立即停止。", 8.42, 4.93, 3.55, 0.3, 12, BLACK)
footer(slide, 2)


# 3. Key contents and status
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = WHITE
heading(slide, "03 / KEY POINTS", "关键技术与项目成果", "汇报时只需要记住下面四个关键词。")

group(slide, 0.72, 1.65, 5.75, 3.95, "关键技术", TEAL)
tech = [
    ("统一入口", "浏览器不直连 ESP32，Go 统一管理"),
    ("控制安全", "权限、租约、限幅、急停、超时"),
    ("视觉闭环", "相机 + YOLO + 路径控制 + 设备反馈"),
    ("可扩展", "WebSocket、WebRTC、OTA、图形化编程"),
]
for i, (h, b) in enumerate(tech):
    y = 2.2 + i * 0.72
    module(slide, 1.02, y, 1.38, 0.42, h, GRAY, 10)
    text(slide, b, 2.62, y + 0.03, 3.35, 0.34, 11, BLACK)

group(slide, 6.85, 1.65, 5.75, 3.95, "当前状态", BLUE)
status = [
    ("自动化测试", "Go / React / Python 已通过", MINT, TEAL),
    ("设备验证", "停止 ACK、认证、重连已验证", PALE_BLUE, BLUE),
    ("继续验收", "断网停机、多鱼并发、公网视频", PALE_ORANGE, ORANGE),
]
for i, (h, b, fill, c) in enumerate(status):
    y = 2.2 + i * 0.86
    module(slide, 7.15, y, 1.48, 0.52, h, fill, 10, c)
    text(slide, b, 8.84, y + 0.05, 3.25, 0.36, 11, BLACK)

rounded(slide, 1.15, 6.0, 11.0, 0.55, GRAY2, BLACK, 0.8)
text(slide, "最终结论：Go 把设备、视觉和用户控制收敛为一条可验证、可恢复、可扩展的链路。", 1.35, 6.13, 10.6, 0.25, 14, TEAL, True, PP_ALIGN.CENTER)
footer(slide, 3)

prs.save(OUT)
print(OUT)
