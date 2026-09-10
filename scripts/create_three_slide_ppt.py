from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from pathlib import Path

OUT = Path("/home/chenfukun/fish/docs/机器鱼技术介绍_2026-09-05.pptx")
W, H = 13.333, 7.5

INK = RGBColor(32, 48, 56)
MUTED = RGBColor(101, 119, 123)
TEAL = RGBColor(36, 157, 137)
BLUE = RGBColor(76, 132, 197)
ORANGE = RGBColor(235, 164, 72)
CORAL = RGBColor(210, 104, 113)
MINT = RGBColor(220, 242, 235)
PAPER = RGBColor(250, 252, 250)
LINE = RGBColor(191, 211, 207)
WHITE = RGBColor(255, 255, 255)

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
blank = prs.slide_layouts[6]


def text(slide, value, x, y, w, h, size=16, color=INK, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def rect(slide, x, y, w, h, fill=WHITE, stroke=INK, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = stroke
    shape.line.width = Pt(1.6)
    if radius:
        shape.adjustments[0] = 0.12
    return shape


def circle(slide, x, y, d, fill=WHITE, stroke=INK):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = stroke
    shape.line.width = Pt(1.8)
    return shape


def arrow(slide, x1, y1, x2, y2, color=TEAL, width=2.0):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    shape.line.end_arrowhead = True
    return shape


def node(slide, x, y, w, h, label, sub, accent, fill):
    rect(slide, x, y, w, h, fill, accent)
    text(slide, label, x, y + 0.18, w, 0.3, 17, INK, True, PP_ALIGN.CENTER)
    text(slide, sub, x + 0.08, y + 0.62, w - 0.16, 0.28, 11, MUTED, False, PP_ALIGN.CENTER)


def header(slide, section, heading, sub):
    text(slide, section, 0.65, 0.4, 3, 0.22, 9, TEAL, True)
    text(slide, heading, 0.65, 0.75, 12, 0.55, 27, INK, True)
    text(slide, sub, 0.68, 1.38, 11.8, 0.3, 12, MUTED)


def footer(slide, number):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(0.65), Inches(7.08), Inches(12.7), Inches(7.08)
    )
    line.line.color.rgb = LINE
    text(slide, "FISH CONTROL SYSTEM", 0.65, 7.16, 2.8, 0.18, 8, MUTED, True)
    text(slide, f"{number:02d}", 12.15, 7.14, 0.55, 0.2, 9, MUTED, True, PP_ALIGN.RIGHT)


# Slide 1: project and topology
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = PAPER
text(slide, "FISH CONTROL SYSTEM", 0.8, 0.55, 4, 0.25, 11, TEAL, True)
text(slide, "机器鱼中央控制系统", 0.8, 1.1, 7.5, 0.6, 32, INK, True)
text(slide, "让浏览器、视觉服务和 ESP32 机器鱼通过一个 Go 中枢协同工作", 0.84, 1.95, 7.3, 0.34, 16, MUTED)

node(slide, 0.9, 3.35, 2.2, 1.2, "浏览器", "React GUI", BLUE, RGBColor(236, 244, 253))
node(slide, 5.35, 3.35, 2.55, 1.2, "Go 控制器", "API / 权限 / 队列", TEAL, MINT)
node(slide, 9.75, 2.55, 2.25, 1.2, "Python 视觉", "相机 / YOLO", ORANGE, RGBColor(255, 246, 226))
node(slide, 9.75, 4.15, 2.25, 1.2, "ESP32 机器鱼", "固件 / 舵机 / OTA", CORAL, RGBColor(253, 237, 239))
arrow(slide, 3.1, 3.95, 5.35, 3.95, BLUE)
arrow(slide, 7.9, 3.78, 9.75, 3.1, ORANGE)
arrow(slide, 7.9, 4.12, 9.75, 4.75, CORAL)
text(slide, "HTTP", 3.55, 3.56, 0.8, 0.22, 10, BLUE, True, PP_ALIGN.CENTER)
text(slide, "内部接口", 8.05, 2.95, 0.95, 0.22, 10, ORANGE, True, PP_ALIGN.CENTER)
text(slide, "WebSocket", 8.05, 4.35, 0.95, 0.22, 10, CORAL, True, PP_ALIGN.CENTER)
text(slide, "一句话：Go 是中枢；Python 看环境；ESP32 执行动作。", 2.25, 6.1, 8.9, 0.34, 17, TEAL, True, PP_ALIGN.CENTER)
footer(slide, 1)


# Slide 2: control and vision loop
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = PAPER
header(slide, "01 / LOOP", "控制闭环：从看见，到运动，再回到状态", "两条链路最终汇合到同一个设备发送队列。")

node(slide, 0.75, 2.2, 2.05, 1.1, "摄像头", "最新帧", BLUE, RGBColor(236, 244, 253))
node(slide, 3.1, 2.2, 2.05, 1.1, "YOLO", "检测 / 跟踪", ORANGE, RGBColor(255, 246, 226))
node(slide, 5.45, 2.2, 2.05, 1.1, "路径控制", "坐标 / PID", TEAL, MINT)
node(slide, 7.8, 2.2, 2.05, 1.1, "Go", "会话 / 租约", CORAL, RGBColor(253, 237, 239))
node(slide, 10.15, 2.2, 2.05, 1.1, "机器鱼", "motion.set", TEAL, WHITE)
for x1, x2, c in [(2.8, 3.1, BLUE), (5.15, 5.45, ORANGE), (7.5, 7.8, TEAL), (9.85, 10.15, CORAL)]:
    arrow(slide, x1, 2.75, x2, 2.75, c)
arrow(slide, 11.2, 3.35, 1.8, 3.35, MUTED, 1.5)
text(slide, "状态反馈", 5.85, 3.48, 1.4, 0.22, 11, MUTED, True, PP_ALIGN.CENTER)

text(slide, "手动控制也走同一条 Go → 队列 → ESP32 链路", 0.95, 4.35, 5.5, 0.3, 15, INK, True)
circle(slide, 1.0, 5.0, 0.38, WHITE, BLUE)
text(slide, "按键", 1.0, 5.08, 0.38, 0.18, 9, BLUE, True, PP_ALIGN.CENTER)
arrow(slide, 1.45, 5.19, 3.3, 5.19, BLUE)
circle(slide, 3.4, 5.0, 0.38, WHITE, TEAL)
text(slide, "Go", 3.4, 5.08, 0.38, 0.18, 9, TEAL, True, PP_ALIGN.CENTER)
arrow(slide, 3.85, 5.19, 5.7, 5.19, TEAL)
circle(slide, 5.8, 5.0, 0.38, WHITE, CORAL)
text(slide, "鱼", 5.8, 5.08, 0.38, 0.18, 9, CORAL, True, PP_ALIGN.CENTER)

text(slide, "关键原则：视觉目标 ID 不等于物理 deviceId，必须显式绑定。", 7.05, 4.92, 5.35, 0.55, 15, CORAL, True)
text(slide, "目标丢失 / 会话过期 / 摄像头失帧 → 停止", 7.15, 5.78, 5.1, 0.28, 13, MUTED)
footer(slide, 2)


# Slide 3: safety, result and next
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = PAPER
header(slide, "02 / SAFETY", "安全与成果：能控制，也知道什么时候必须停止", "把当前完成度和下一步边界放在同一张图里。")

guards = [("心跳", "约 3 秒", TEAL), ("视觉超时", "约 3 秒", CORAL),
          ("控制权", "账号 + 窗口", BLUE), ("参数限幅", "中位 ± 振幅", ORANGE)]
for i, (h, b, c) in enumerate(guards):
    x = 0.82 + i * 3.05
    circle(slide, x + 0.78, 2.05, 0.9, WHITE, c)
    text(slide, "停", x + 0.78, 2.3, 0.9, 0.25, 17, c, True, PP_ALIGN.CENTER)
    text(slide, h, x, 3.2, 2.45, 0.25, 15, INK, True, PP_ALIGN.CENTER)
    text(slide, b, x, 3.58, 2.45, 0.25, 11, MUTED, False, PP_ALIGN.CENTER)
    if i < 3:
        arrow(slide, x + 1.72, 2.5, x + 2.75, 2.5, LINE, 1.5)

rect(slide, 0.88, 4.45, 5.45, 1.22, MINT, TEAL)
text(slide, "已完成", 1.15, 4.7, 1.0, 0.25, 16, TEAL, True)
text(slide, "Go / React / Python 自动化测试通过\n设备停止 ACK、认证、租约、视觉超时已验证", 2.25, 4.62, 3.7, 0.55, 12, INK)

rect(slide, 7.0, 4.45, 5.45, 1.22, RGBColor(253, 237, 239), CORAL)
text(slide, "下一步", 7.28, 4.7, 1.0, 0.25, 16, CORAL, True)
text(slide, "实机停机时延、断网恢复、多鱼并发\n公网视频稳定性和 Mixly 程序执行", 8.38, 4.62, 3.7, 0.55, 12, INK)

text(slide, "演示顺序：发现设备 → 领取 → 手动控制 → 视觉绑定 → 自动停止保护", 1.25, 6.25, 10.9, 0.3, 15, TEAL, True, PP_ALIGN.CENTER)
footer(slide, 3)

prs.save(OUT)
print(OUT)
