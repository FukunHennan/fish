from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path("/home/chenfukun/fish/docs/机器鱼技术介绍_2026-09-05.pptx")
W, H = 13.333, 7.5

BLACK = RGBColor(28, 35, 38)
GRAY = RGBColor(232, 232, 232)
GRAY2 = RGBColor(245, 245, 245)
GRAY3 = RGBColor(205, 205, 205)
WHITE = RGBColor(255, 255, 255)
TEAL = RGBColor(33, 137, 119)
BLUE = RGBColor(61, 106, 158)
ORANGE = RGBColor(178, 115, 35)
RED = RGBColor(177, 70, 79)

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
blank = prs.slide_layouts[6]


def shape(slide, x, y, w, h, fill=WHITE, stroke=BLACK, radius=True, lw=1.0):
    s = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = stroke
    s.line.width = Pt(lw)
    if radius:
        s.adjustments[0] = 0.08
    return s


def text(slide, value, x, y, w, h, size=12, color=BLACK, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    t = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = t.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = value
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return t


def connector(slide, x1, y1, x2, y2, color=BLACK, width=1.2, arrow=False):
    c = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if arrow:
        c.line.end_arrowhead = True
    return c


def outer(slide, x, y, w, h, label, accent=BLACK):
    shape(slide, x, y, w, h, WHITE, BLACK, True, 1.2)
    text(slide, label, x + 0.12, y + 0.06, w - 0.24, 0.25, 13, accent, True, PP_ALIGN.CENTER)


def module(slide, x, y, w, h, label, fill=GRAY, size=11, color=BLACK):
    shape(slide, x, y, w, h, fill, BLACK, True, 0.8)
    text(slide, label, x + 0.04, y + 0.03, w - 0.08, h - 0.06, size, color, False, PP_ALIGN.CENTER)


def footer(slide, n):
    connector(slide, 0.55, 7.08, 12.78, 7.08, GRAY3, 0.7)
    text(slide, "FISH CONTROL SYSTEM", 0.58, 7.14, 2.5, 0.18, 8, GRAY3, True, PP_ALIGN.LEFT)
    text(slide, f"{n:02d}", 12.15, 7.13, 0.55, 0.2, 9, GRAY3, True, PP_ALIGN.RIGHT)


def heading(slide, kicker, title, sub):
    text(slide, kicker, 0.62, 0.28, 3.0, 0.22, 9, TEAL, True, PP_ALIGN.LEFT)
    text(slide, title, 0.62, 0.56, 11.9, 0.42, 23, BLACK, True, PP_ALIGN.LEFT)
    text(slide, sub, 0.64, 1.04, 11.8, 0.25, 11, GRAY3, False, PP_ALIGN.LEFT)


# Slide 1: nested block diagram, matching the reference style.
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = WHITE
heading(slide, "01 / SYSTEM BLOCK DIAGRAM", "机器鱼中央控制系统", "浏览器、Go、Python 和 ESP32 的职责边界")

shape(slide, 0.5, 1.48, 12.33, 5.25, WHITE, BLACK, True, 1.3)
text(slide, "机器鱼中央控制系统 / Fish Control System", 0.72, 1.57, 4.5, 0.22, 12, BLACK, True)

outer(slide, 0.72, 1.95, 5.82, 1.82, "用户与前端", BLUE)
module(slide, 0.94, 2.38, 1.25, 0.52, "React GUI")
module(slide, 2.35, 2.38, 1.25, 0.52, "设备列表")
module(slide, 3.76, 2.38, 1.25, 0.52, "手动控制")
module(slide, 5.17, 2.38, 1.15, 0.52, "视觉工作区")
module(slide, 1.65, 3.08, 1.25, 0.42, "账号登录", GRAY2, 10)
module(slide, 3.06, 3.08, 1.25, 0.42, "多用户租约", GRAY2, 10)
module(slide, 4.47, 3.08, 1.25, 0.42, "OTA 页面", GRAY2, 10)

outer(slide, 6.8, 1.95, 5.8, 1.82, "Go 中央控制器", TEAL)
module(slide, 7.03, 2.38, 1.32, 0.52, "HTTP API", MINT := RGBColor(224, 242, 236))
module(slide, 8.52, 2.38, 1.32, 0.52, "Auth / Session")
module(slide, 10.01, 2.38, 1.32, 0.52, "Device Hub")
module(slide, 11.5, 2.38, 0.85, 0.52, "队列")
module(slide, 7.85, 3.08, 1.32, 0.42, "Vision Proxy", GRAY2, 10)
module(slide, 9.34, 3.08, 1.32, 0.42, "OTA Manager", GRAY2, 10)
module(slide, 10.83, 3.08, 1.32, 0.42, "诊断日志", GRAY2, 10)

outer(slide, 0.72, 4.02, 7.05, 2.35, "Python 视觉服务", ORANGE)
module(slide, 0.95, 4.45, 1.4, 0.58, "Camera")
module(slide, 2.55, 4.45, 1.4, 0.58, "YOLO")
module(slide, 4.15, 4.45, 1.4, 0.58, "Tracking")
module(slide, 5.75, 4.45, 1.4, 0.58, "WebRTC")
module(slide, 1.75, 5.28, 1.4, 0.52, "标定 / 坐标", GRAY2, 10)
module(slide, 3.35, 5.28, 1.4, 0.52, "路径 / 航向", GRAY2, 10)
module(slide, 4.95, 5.28, 1.4, 0.52, "PID / 动作", GRAY2, 10)

outer(slide, 8.0, 4.02, 4.6, 2.35, "ESP32 机器鱼", RED)
module(slide, 8.24, 4.45, 1.23, 0.58, "Discovery")
module(slide, 9.65, 4.45, 1.23, 0.58, "HMAC")
module(slide, 11.06, 4.45, 1.23, 0.58, "WebSocket")
module(slide, 8.95, 5.28, 1.23, 0.52, "运动执行", GRAY2, 10)
module(slide, 10.36, 5.28, 1.23, 0.52, "心跳状态", GRAY2, 10)
module(slide, 11.77, 5.28, 0.58, 0.52, "OTA", GRAY2, 10)

connector(slide, 6.54, 2.85, 6.8, 2.85, BLUE, 1.8)
connector(slide, 9.68, 3.77, 9.68, 4.02, TEAL, 1.8)
connector(slide, 7.77, 5.18, 8.0, 5.18, ORANGE, 1.8)
text(slide, "统一入口", 6.55, 3.0, 1.2, 0.2, 9, BLUE, True, PP_ALIGN.CENTER)
text(slide, "内部 HTTP", 9.1, 3.78, 1.2, 0.2, 9, TEAL, True, PP_ALIGN.CENTER)
text(slide, "设备 WebSocket", 7.55, 5.28, 1.65, 0.2, 9, RED, True, PP_ALIGN.CENTER)
footer(slide, 1)


# Slide 2: internal block topology.
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = WHITE
heading(slide, "02 / DATA PATH", "一条运动命令经过哪些模块", "把输入、判断、执行和反馈画成模块链")

shape(slide, 0.62, 1.55, 12.1, 4.95, WHITE, BLACK, True, 1.3)
text(slide, "控制面", 0.85, 1.72, 1.1, 0.22, 12, BLUE, True)
text(slide, "浏览器", 1.0, 2.15, 1.0, 0.22, 12, BLACK, True, PP_ALIGN.CENTER)
module(slide, 0.78, 2.52, 1.45, 0.62, "按键 / 视觉意图")
module(slide, 2.65, 2.52, 1.45, 0.62, "用户权限")
module(slide, 4.52, 2.52, 1.45, 0.62, "设备租约")
module(slide, 6.39, 2.52, 1.45, 0.62, "参数限幅")
module(slide, 8.26, 2.52, 1.45, 0.62, "单鱼发送队列")
module(slide, 10.13, 2.52, 1.45, 0.62, "ESP32 执行")
for x in [2.23, 4.1, 5.97, 7.84, 9.71]:
    connector(slide, x, 2.83, x + 0.42, 2.83, BLACK, 1.4)
text(slide, "↓", 6.36, 3.42, 0.25, 0.3, 19, TEAL, True, PP_ALIGN.CENTER)
text(slide, "状态面", 0.85, 3.83, 1.1, 0.22, 12, TEAL, True)
module(slide, 1.25, 4.32, 1.8, 0.62, "heartbeat / state", MINT)
module(slide, 3.75, 4.32, 1.8, 0.62, "command.result")
module(slide, 6.25, 4.32, 1.8, 0.62, "Go 汇总状态")
module(slide, 8.75, 4.32, 1.8, 0.62, "GUI 更新")
connector(slide, 3.05, 4.63, 3.75, 4.63, TEAL, 1.4, True)
connector(slide, 5.55, 4.63, 6.25, 4.63, TEAL, 1.4, True)
connector(slide, 8.05, 4.63, 8.75, 4.63, TEAL, 1.4, True)

shape(slide, 1.1, 5.55, 10.9, 0.57, GRAY2, BLACK, True, 0.8)
text(slide, "关键语义：queued 只是入队；sent 是写入设备；acknowledged 才是设备返回结果。", 1.25, 5.68, 10.6, 0.25, 12, RED, True, PP_ALIGN.CENTER)
footer(slide, 2)


# Slide 3: safety and completion blocks.
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = WHITE
heading(slide, "03 / SAFETY & STATUS", "安全机制与当前完成度", "用模块框表达“已实现”和“待验收”")

outer(slide, 0.7, 1.62, 6.05, 4.7, "安全模块", RED)
module(slide, 0.98, 2.1, 1.45, 0.62, "设备心跳", RGBColor(253, 237, 239))
module(slide, 2.7, 2.1, 1.45, 0.62, "视觉超时", RGBColor(253, 237, 239))
module(slide, 4.42, 2.1, 1.45, 0.62, "控制权", RGBColor(253, 237, 239))
module(slide, 1.84, 3.05, 1.45, 0.62, "参数限幅", GRAY2)
module(slide, 3.56, 3.05, 1.45, 0.62, "急停", GRAY2)
module(slide, 1.0, 4.0, 1.45, 0.62, "目标丢失", GRAY2)
module(slide, 2.72, 4.0, 1.45, 0.62, "摄像头失帧", GRAY2)
module(slide, 4.44, 4.0, 1.45, 0.62, "断线停止", GRAY2)
shape(slide, 1.0, 5.15, 5.45, 0.62, GRAY2, BLACK, True, 0.8)
text(slide, "任一条件失效 → 停止命令进入同一设备队列", 1.18, 5.31, 5.1, 0.25, 11, RED, True, PP_ALIGN.CENTER)

outer(slide, 7.0, 1.62, 5.62, 4.7, "项目状态", TEAL)
module(slide, 7.3, 2.1, 1.55, 0.72, "Go", MINT, 13)
module(slide, 9.1, 2.1, 1.55, 0.72, "React", MINT, 13)
module(slide, 10.9, 2.1, 1.4, 0.72, "Python", MINT, 13)
text(slide, "自动化测试通过", 7.52, 3.0, 4.45, 0.27, 14, TEAL, True, PP_ALIGN.CENTER)
module(slide, 7.3, 3.65, 1.55, 0.72, "设备", GRAY2, 13)
module(slide, 9.1, 3.65, 1.55, 0.72, "OTA", GRAY2, 13)
module(slide, 10.9, 3.65, 1.4, 0.72, "视觉", GRAY2, 13)
text(slide, "部分实机验证", 7.52, 4.56, 4.45, 0.27, 14, ORANGE, True, PP_ALIGN.CENTER)
shape(slide, 7.3, 5.15, 4.99, 0.62, GRAY2, BLACK, True, 0.8)
text(slide, "下一步：断网停机、多鱼并发、公网视频", 7.45, 5.31, 4.7, 0.25, 11, ORANGE, True, PP_ALIGN.CENTER)

footer(slide, 3)

prs.save(OUT)
print(OUT)
